"""xlsx lineart シートの読み込み。

元GAS: sheet.js の getDic(SSID, sheetName).arrayByRow() に相当。
1行目=ヘッダ（列名）、2行目=section行（表示用の見出し。データではない）、
3行目以降が実データ行という原稿ブックのレイアウトにならう。
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

Row = dict[str, object]

# メタ列（カテゴリ集計の対象にしない列）
META_COLUMNS = {"page", "plot", "mode", "scene", "title"}


def load_sheet(path: str | Path, sheet_name: str = "lineart") -> tuple[list[Row], list[str]]:
    """xlsx から (データ行のリスト, カテゴリ列名のリスト) を返す。

    データ行は {列名: セル値} の dict。空セルは None のまま保持する
    （data.js の `work[k] == ''` 判定に対応するのは analyze.py 側で行う）。
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"入力ファイルが見つかりません: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"シート '{sheet_name}' が見つかりません: {path} "
            f"(存在するシート: {', '.join(wb.sheetnames)})"
        )
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)

    header_raw = next(rows_iter, None)
    if header_raw is None:
        raise ValueError(f"シート '{sheet_name}' にヘッダ行がありません")
    header = [str(h).strip() if h is not None else "" for h in header_raw]

    if "plot" not in header or "mode" not in header:
        raise ValueError(
            f"シート '{sheet_name}' のヘッダに 'plot' または 'mode' 列がありません: {header}"
        )

    next(rows_iter, None)  # 2行目 = section行（表示用の見出し）をスキップ

    categories = [h for h in header if h and h not in META_COLUMNS]

    rows: list[Row] = []
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        row: Row = {}
        for i, name in enumerate(header):
            if not name:
                continue
            row[name] = raw[i] if i < len(raw) else None
        rows.append(row)

    return rows, categories
